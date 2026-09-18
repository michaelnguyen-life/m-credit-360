"""Generate MSB-formal enterprise credit memo (MB02a/QT.RR.037) in DOCX format."""

from __future__ import annotations

import json
import logging
import os
import re
from datetime import datetime, timezone
from decimal import Decimal
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from eb_credit_agent import AssessmentError, CreditAssessment, _decimal


LOGGER = logging.getLogger(__name__)
MSB_RED = "C41230"
ALERT_YELLOW = "FFF3CD"
ALERT_TEXT = "856404"
GRAY = "B7B7B7"
NOT_PROVIDED = "Ch\u01b0a cung c\u1ea5p"
FORM_CODE = "MB02a/QT.RR.037"

FINANCIAL_ROWS = (
    ("Doanh thu thu\u1ea7n", ("IS_REVENUE", "REVENUE")),
    ("T\u1ed5ng t\u00e0i s\u1ea3n", ("BS_TOTAL_ASSETS", "TOTAL_ASSETS")),
    ("N\u1ee3 ng\u1eafn h\u1ea1n", ("BS_CURRENT_LIABILITIES", "CURRENT_LIABILITIES")),
    ("N\u1ee3 d\u00e0i h\u1ea1n", ("BS_LONG_TERM_DEBT", "LONG_TERM_DEBT")),
    ("V\u1ed1n ch\u1ee7 s\u1edf h\u1eefu", ("BS_EQUITY", "EQUITY")),
    ("L\u1ee3i nhu\u1eadn g\u1ed9p", ("IS_GROSS_PROFIT", "GROSS_PROFIT")),
    ("EBIT", ("IS_EBIT", "EBIT")),
    ("Chi ph\u00ed l\u00e3i vay", ("IS_INTEREST_EXPENSE", "INTEREST_EXPENSE")),
    ("L\u1ee3i nhu\u1eadn sau thu\u1ebf", ("IS_NET_PROFIT", "NET_PROFIT")),
    ("D\u00f2ng ti\u1ec1n H\u0110KD (CFO)", ("CF_OPERATING_CASH_FLOW", "CFO")),
)


class MemoBuilderError(ValueError):
    """A Word memo request is invalid or cannot be rendered."""


def _money(value: Any) -> str:
    numeric = _decimal(value)
    if numeric is None:
        return NOT_PROVIDED
    return f"{int(numeric):,}".replace(",", ".") + " VND"


def _ratio(value: Any) -> str:
    numeric = _decimal(value)
    return f"{float(numeric):.2f}x" if numeric is not None else NOT_PROVIDED


def _percent(value: Any) -> str:
    numeric = _decimal(value)
    return f"{float(numeric) * 100:.2f}%" if numeric is not None else NOT_PROVIDED


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _cell_border(cell: Any, color: str = GRAY) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


class CreditMemoBuilder:
    """Creates a DOCX memo from an EB payload or a completed Credit Profile."""

    def __init__(self, output_dir: str | Path = "output_memos") -> None:
        self.output_dir = Path(output_dir)
        self.credit_agent = CreditAssessment()

    def build(self, request: Mapping[str, Any]) -> dict[str, Any]:
        if not isinstance(request, Mapping):
            raise MemoBuilderError("Request must be a JSON object.")
        profile, source_payload = self._resolve_profile(request)
        warnings = list(profile.get("warnings", []))
        document = self._new_document()
        company = profile.get("company", {})
        report_year = self._report_year(profile.get("reporting_period"))
        memo_number = str(request.get("memo_number") or f"MSB/TTTD/{report_year}/{profile['assessment_id']}")
        self._header(document, memo_number)
        self._part1_customer_legal(document, company, source_payload)
        self._part2_credit_proposal(document, profile, source_payload)
        self._part3_appraisal_financials(document, profile, source_payload, warnings)
        self._part4_supply_chain(document, profile)
        self._part5_statement_anomaly(document, profile)
        self._part6_product_039(document, profile)
        self._part7_swot(document, profile)
        self._part8_conclusion_signatures(document, profile, source_payload, warnings)
        self._footer(document)
        tax_id = re.sub(r"\D", "", str(company.get("tax_id") or "UNKNOWN")) or "UNKNOWN"
        self.output_dir.mkdir(parents=True, exist_ok=True)
        output_path = self.output_dir / f"TO_TRINH_TIN_DUNG_MB02a_{tax_id}_{report_year}.docx"
        document.save(output_path)
        return {"status": "completed_with_warnings" if warnings else "completed", "file_path": str(output_path), "assessment_id": profile["assessment_id"], "preliminary_decision": profile.get("preliminary_decision", {}).get("outcome"), "product_039_decision": profile.get("product_039_evaluation", {}).get("decision"), "warnings": warnings}

    def _resolve_profile(self, request: Mapping[str, Any]) -> tuple[dict[str, Any], Mapping[str, Any]]:
        if isinstance(request.get("credit_profile"), Mapping):
            source = request.get("eb_payload", {})
            return dict(request["credit_profile"]), source if isinstance(source, Mapping) else {}
        payload = request.get("eb_payload", request)
        if not isinstance(payload, Mapping):
            raise MemoBuilderError("Provide credit_profile or eb_payload as a JSON object.")
        try:
            return self.credit_agent.assess(payload), payload
        except AssessmentError as exc:
            raise MemoBuilderError(str(exc)) from exc

    @staticmethod
    def _report_year(period: Any) -> str:
        found = re.search(r"(20\d{2})", str(period or ""))
        return found.group(1) if found else str(datetime.now(timezone.utc).year)

    @staticmethod
    def _new_document() -> Document:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.line_spacing = 1.15
        return document

    def _header(self, document: Document, memo_number: str) -> None:
        p1 = document.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run("NG\u00c2N H\u00c0NG TMCP H\u00c0NG H\u1ea2I VI\u1ec6T NAM (MSB)")
        r1.bold = True; r1.font.name = "Arial"; r1.font.size = Pt(12)
        r1.font.color.rgb = RGBColor(196, 18, 48)
        p2 = document.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"M\u00e3 bi\u1ec3u m\u1eabu: {FORM_CODE}")
        r2.font.size = Pt(10); r2.italic = True
        p3 = document.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run("T\u1edc TR\u00ccNH TH\u1ea8M \u0110\u1ecaNH & PH\u00ca DUY\u1ec6T T\u00cdN D\u1ee4NG DOANH NGHI\u1ec6P")
        r3.bold = True; r3.font.size = Pt(14)
        r3.font.color.rgb = RGBColor(196, 18, 48)
        p4 = document.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run(f"S\u1ed1 t\u1edd tr\u00ecnh: {memo_number}  |  Ng\u00e0y l\u1eadp: {datetime.now().strftime('%d/%m/%Y')}")
        r4.font.size = Pt(10)

    def _heading(self, document: Document, text: str) -> None:
        p = document.add_heading(text, level=1)
        for r in p.runs:
            r.font.color.rgb = RGBColor(196, 18, 48); r.font.name = "Arial"

    def _subheading(self, document: Document, text: str) -> None:
        p = document.add_heading(text, level=2)
        for r in p.runs:
            r.font.color.rgb = RGBColor(60, 60, 60); r.font.name = "Arial"

    def _part1_customer_legal(self, document: Document, company: Mapping[str, Any], source: Mapping[str, Any]) -> None:
        self._heading(document, "I. TH\u00d4NG TIN KH\u00c1CH H\u00c0NG & PH\u00c1P L\u00dd")
        legal = source.get("legal", {}) if isinstance(source.get("legal"), Mapping) else {}
        cs = source.get("company", {}) if isinstance(source.get("company"), Mapping) else {}
        name = company.get("name") or cs.get("name") or NOT_PROVIDED
        tax_id = company.get("tax_id") or cs.get("tax_id") or NOT_PROVIDED
        cif = company.get("cif") or cs.get("cif") or NOT_PROVIDED
        c_status = company.get("customer_status") or cs.get("customer_status") or "M\u1edbi"
        c_segment = company.get("customer_segment") or cs.get("customer_segment") or "SME"
        reg_num = legal.get("registration_number") or cs.get("registration_number") or NOT_PROVIDED
        reg_date = legal.get("registration_date") or cs.get("registration_date") or NOT_PROVIDED
        reg_place = legal.get("registration_place") or cs.get("registration_place") or NOT_PROVIDED
        op_years = company.get("operating_years") or cs.get("operating_years") or NOT_PROVIDED
        charter = company.get("charter_capital") or cs.get("charter_capital") or legal.get("charter_capital")
        paid = legal.get("paid_capital") or cs.get("paid_capital")
        industry = company.get("industry") or cs.get("industry") or NOT_PROVIDED
        ind_l3 = legal.get("industry_level3") or cs.get("industry_level3") or industry
        ind_l5 = legal.get("industry_level5") or cs.get("industry_level5") or NOT_PROVIDED
        risk_sec = legal.get("risk_sector") or cs.get("risk_sector") or NOT_PROVIDED
        addr = company.get("registered_address") or cs.get("registered_address") or cs.get("address") or NOT_PROVIDED
        rep = company.get("legal_rep") or cs.get("legal_rep") or NOT_PROVIDED
        rep_id = legal.get("legal_rep_id") or cs.get("legal_rep_id") or cs.get("cccd") or NOT_PROVIDED
        rep_id_date = legal.get("legal_rep_id_date") or cs.get("legal_rep_id_date") or NOT_PROVIDED
        table = self._table(document, ["H\u1ea1ng m\u1ee5c", "N\u1ed9i dung", "H\u1ea1ng m\u1ee5c", "N\u1ed9i dung"])
        rows = [
            ("T\u00ean Doanh nghi\u1ec7p", name, "M\u00e3 CIF", cif),
            ("M\u00e3 s\u1ed1 thu\u1ebf (MST)", tax_id, "T\u00ecnh tr\u1ea1ng KH", c_status),
            ("\u0110\u1ed1i t\u01b0\u1ee3ng KH", c_segment, "S\u1ed1 \u0110KKD", reg_num),
            ("Ng\u00e0y c\u1ea5p \u0110KKD", reg_date, "N\u01a1i c\u1ea5p \u0110KKD", reg_place),
            ("S\u1ed1 n\u0103m ho\u1ea1t \u0111\u1ed9ng", op_years, "V\u1ed1n \u0111i\u1ec1u l\u1ec7 \u0111\u0103ng k\u00fd", _money(charter) if charter else NOT_PROVIDED),
            ("V\u1ed1n \u0111i\u1ec1u l\u1ec7 th\u1ef1c g\u00f3p", _money(paid) if paid else NOT_PROVIDED, "Ng\u00e0nh ngh\u1ec1 c\u1ea5p 3", ind_l3),
            ("Ng\u00e0nh ngh\u1ec1 c\u1ea5p 5", ind_l5, "Nh\u00f3m ng\u00e0nh r\u1ee7i ro", risk_sec),
            ("\u0110\u1ecba ch\u1ec9 tr\u1ee5 s\u1edf", addr, "", ""),
            ("Ng\u01b0\u1eddi \u0111\u1ea1i di\u1ec7n PL", rep, "S\u1ed1 CCCD", rep_id),
            ("Ng\u00e0y c\u1ea5p CCCD", rep_id_date, "", ""),
        ]
        for row in rows:
            self._row(table, list(row))

    def _part2_credit_proposal(self, document: Document, profile: Mapping[str, Any], source: Mapping[str, Any]) -> None:
        self._heading(document, "II. \u0110\u1ec0 XU\u1ea4T C\u1ea4P T\u00cdN D\u1ee4NG & T\u00c0I S\u1ea2N B\u1ea2O \u0110\u1ea2M")
        ct = source.get("credit_terms", {}) if isinstance(source.get("credit_terms"), Mapping) else {}
        pc = profile.get("product_039_evaluation", {}).get("calculations", {})
        metrics = profile.get("metrics", {})
        revenue = metrics.get("revenue")
        self._subheading(document, "1. \u0110\u1ec1 xu\u1ea5t H\u1ea1n m\u1ee9c T\u00edn d\u1ee5ng (B5)")
        t = self._table(document, ["H\u1ea1ng m\u1ee5c", "H\u1ea1n m\u1ee9c (VND)", "Th\u1eddi h\u1ea1n", "Ghi ch\u00fa"])
        stl = ct.get("short_term_limit") or pc.get("loan_request_amount")
        tfl = ct.get("trade_finance_limit")
        gl = ct.get("guarantee_limit")
        tenor = ct.get("tenor") or "12 th\u00e1ng"
        self._row(t, ["Cho vay ng\u1eafn h\u1ea1n (Main 1)", _money(stl), tenor, ""])
        self._row(t, ["T\u00e0i tr\u1ee3 TM L/C (Main 2)", _money(tfl), tenor, ""])
        self._row(t, ["B\u1ea3o l\u00e3nh ngo\u1ea1i b\u1ea3ng", _money(gl), tenor, ""])
        mcl_vals = [v for v in [stl, tfl, gl] if v is not None]
        mcl = sum(_decimal(v) for v in mcl_vals if _decimal(v) is not None) if mcl_vals else None
        self._row(t, ["T\u1ed5ng gi\u1edbi h\u1ea1n HMTD (MCL)", _money(mcl), "", ""])
        lr = _decimal(mcl) / _decimal(revenue) if mcl and revenue and _decimal(revenue) else None
        self._row(t, ["T\u1ef7 l\u1ec7 HMTD/Doanh thu", _percent(lr), "", ""])
        self._subheading(document, "2. T\u00e0i s\u1ea3n B\u1ea3o \u0111\u1ea3m (B7)")
        coll = source.get("collateral", []) if isinstance(source.get("collateral"), list) else []
        ct2 = self._table(document, ["Lo\u1ea1i TS", "T\u00ean t\u00e0i s\u1ea3n", "Ch\u1ee7 TS", "M\u1ed1i QH", "Gi\u00e1 tr\u1ecb \u0111\u1ecbnh gi\u00e1", "T\u1ef7 l\u1ec7 K"])
        if coll:
            for item in coll:
                if isinstance(item, Mapping):
                    self._row(ct2, [str(item.get("type", NOT_PROVIDED)), str(item.get("name", NOT_PROVIDED)), str(item.get("owner", NOT_PROVIDED)), str(item.get("relationship", NOT_PROVIDED)), _money(item.get("value")), _percent(item.get("allocation_ratio")) if item.get("allocation_ratio") else NOT_PROVIDED])
        else:
            self._row(ct2, ["", str(ct.get("collateral", NOT_PROVIDED)), "", "", "", ""])
        self._subheading(document, "3. \u0110i\u1ec1u ki\u1ec7n T\u00edn d\u1ee5ng (B13)")
        ct3 = self._table(document, ["\u0110i\u1ec1u ki\u1ec7n", "N\u1ed9i dung", "Ghi ch\u00fa"])
        self._row(ct3, ["Th\u1eddi h\u1ea1n duy tr\u00ec HMTD", tenor, ""])
        self._row(ct3, ["L\u00e3i su\u1ea5t cho vay", str(ct.get("interest_rate", NOT_PROVIDED)), ""])
        self._row(ct3, ["T\u1ef7 l\u1ec7 k\u00fd qu\u1ef9 t\u1ed1i thi\u1ec3u", str(ct.get("deposit_ratio", NOT_PROVIDED)), ""])

    def _part3_appraisal_financials(self, document: Document, profile: Mapping[str, Any], source: Mapping[str, Any], warnings: list[str]) -> None:
        self._heading(document, "III. N\u1ed8I DUNG TH\u1ea8M \u0110\u1ecaNH & T\u00c0I CH\u00cdNH BCTC 3 K\u1ef2")
        self._subheading(document, "1. Quan h\u1ec7 t\u00edn d\u1ee5ng & CIC (B16)")
        cic = source.get("cic", {}) if isinstance(source.get("cic"), Mapping) else {}
        t1 = self._table(document, ["Ch\u1ec9 ti\u00eau", "T\u1ea1i MSB", "T\u1ea1i TCTD kh\u00e1c", "T\u1ed5ng", "Ghi ch\u00fa"])
        self._row(t1, ["D\u01b0 n\u1ee3 ng\u1eafn h\u1ea1n", _money(cic.get("msb_short_term_debt")), _money(cic.get("other_short_term_debt")), "", ""])
        self._row(t1, ["D\u01b0 n\u1ee3 TD d\u00e0i h\u1ea1n", _money(cic.get("msb_long_term_debt")), _money(cic.get("other_long_term_debt")), "", ""])
        self._row(t1, ["Nh\u00f3m n\u1ee3 CIC", str(cic.get("debt_group", "Nh\u00f3m 1")), "", "", ""])
        self._row(t1, ["L\u1ecbch s\u1eed tr\u1ea3 n\u1ee3", str(cic.get("payment_history", "T\u1ed1t")), "", "", ""])
        self._subheading(document, "2. T\u1ed5ng h\u1ee3p BCTC 3 n\u0103m (B32)")
        history = source.get("financial_history") if isinstance(source.get("financial_history"), list) else []
        current = {"period": profile.get("reporting_period"), "financials": profile.get("normalized_financials", {})}
        periods = (history + [current])[-3:]
        headers = ["Ch\u1ec9 ti\u00eau"] + [str(item.get("period") or NOT_PROVIDED) for item in periods]
        while len(headers) < 4:
            headers.insert(1, NOT_PROVIDED)
        t2 = self._table(document, headers)
        for label, codes in FINANCIAL_ROWS:
            cells = [label]
            for item in periods:
                fin = item.get("financials", {}) if isinstance(item, Mapping) else {}
                value = next((fin.get(code) for code in codes if isinstance(fin, Mapping) and code in fin), None)
                cells.append(_money(value))
            while len(cells) < 4:
                cells.insert(1, NOT_PROVIDED)
            self._row(t2, cells)
        self._subheading(document, "3. Ch\u1ec9 s\u1ed1 an to\u00e0n t\u00e0i ch\u00ednh & D\u00f2ng ti\u1ec1n (B33)")
        ratios = profile.get("ratios", {})
        t3 = self._table(document, ["Ch\u1ec9 s\u1ed1", "Gi\u00e1 tr\u1ecb", "Ng\u01b0\u1edfng", "Nh\u1eadn x\u00e9t"])
        nwc_val = _decimal(ratios.get("nwc"))
        dscr_val = _decimal(ratios.get("dscr"))
        icr_val = _decimal(ratios.get("icr"))
        self._row(t3, ["NWC", _money(ratios.get("nwc")), "> 0", "D\u01b0\u01a1ng" if nwc_val is not None and nwc_val >= 0 else "\u00c2m"], alert=nwc_val is not None and nwc_val < 0)
        self._row(t3, ["WCR", _money(ratios.get("wcr")), "\u2014", "Theo d\u00f5i"])
        self._row(t3, ["DSCR", _ratio(ratios.get("dscr")), "\u2265 1,00x", "\u0110\u1ea1t" if dscr_val is not None and dscr_val >= 1 else "D\u01b0\u1edbi ng\u01b0\u1edfng"], alert=dscr_val is not None and dscr_val < 1)
        self._row(t3, ["ICR", _ratio(ratios.get("icr")), "\u2265 1,50x", "\u0110\u1ea1t" if icr_val is not None and icr_val >= Decimal("1.5") else "D\u01b0\u1edbi ng\u01b0\u1edfng"], alert=icr_val is not None and icr_val < Decimal("1.5"))
        self._row(t3, ["D\u00f2ng ti\u1ec1n th\u1ee7an", _money(ratios.get("net_cash_flow")), "\u2014", "Theo BCTC"])
        self._subheading(document, "4. \u0110\u1ed1i so\u00e1t Digisale (DSP) & 5 Red Flags")
        dsp = source.get("dsp", {}) if isinstance(source.get("dsp"), Mapping) else {}
        metrics = profile.get("metrics", {})
        t4 = self._table(document, ["Ch\u1ec9 ti\u00eau", "BCTC", "DSP", "Ch\u00eanh l\u1ec7ch", "T\u1ef7 l\u1ec7 l\u1ec7ch"])
        for label, fkey, dkey in (("Doanh thu", "revenue", "revenue"), ("Ph\u1ea3i thu TM", "trade_receivables", "trade_receivables")):
            bctc, dv = _decimal(metrics.get(fkey)), _decimal(dsp.get(dkey))
            diff = abs(bctc - dv) if bctc is not None and dv is not None else None
            var = diff / abs(bctc) if diff is not None and bctc else None
            self._row(t4, [label, _money(bctc), _money(dv), _money(diff), _percent(var) if var is not None else NOT_PROVIDED], alert=var is not None and var > Decimal("0.05"))
        t5 = self._table(document, ["M\u00e3", "R\u1ee7i ro", "Tr\u1ea1ng th\u00e1i", "M\u1ee9c \u0111\u1ed9", "B\u1eb1ng ch\u1ee9ng"])
        for flag in profile.get("red_flags", []):
            status = "RED FLAG" if flag.get("triggered") else "PASS"
            self._row(t5, [str(flag.get("code")), str(flag.get("title")), status, str(flag.get("severity", "none")).upper(), str(flag.get("evidence"))], alert=bool(flag.get("triggered")))
        self._subheading(document, "5. Ph\u00e2n t\u00edch chuy\u00ean s\u00e2u s\u1ee9c kh\u1ecfe t\u00e0i ch\u00ednh")
        deep_dive = profile.get("financial_deep_dive", {})
        t6 = self._table(document, ["L\u0129nh v\u1ef1c", "Quan s\u00e1t", "\u0110\u00e1nh gi\u00e1"])
        for item in deep_dive.get("assessments", []):
            self._row(t6, [str(item.get("area", NOT_PROVIDED)), str(item.get("observation", NOT_PROVIDED)), str(item.get("assessment", NOT_PROVIDED))])

    def _part4_supply_chain(self, document: Document, profile: Mapping[str, Any]) -> None:
        self._heading(document, "IV. PH\u00c2N T\u00cdCH CHU\u1ed4I CUNG \u1ee8NG (TOP 5)")
        sc = profile.get("supply_chain_analysis", {})
        self._subheading(document, "1. Top 5 Nh\u00e0 cung c\u1ea5p (Suppliers)")
        t1 = self._table(document, ["#", "T\u00ean nh\u00e0 cung c\u1ea5p", "Gi\u00e1 tr\u1ecb mua", "T\u1ef7 l\u1ec7", "Terms", "S\u1ed1 n\u0103m QH"])
        for i, supplier in enumerate(sc.get("top_suppliers", []), 1):
            self._row(t1, [str(i), str(supplier.get("name", NOT_PROVIDED)), _money(supplier.get("amount")), _percent(supplier.get("ratio")), str(supplier.get("payment_terms", NOT_PROVIDED)), str(supplier.get("relationship_years", NOT_PROVIDED))])
        p1 = document.add_paragraph()
        r1 = p1.add_run(f"T\u1ed5ng t\u1ef7 l\u1ec7 Top 5: {_percent(sc.get('top5_supplier_ratio'))}  |  M\u1ee9c \u0111\u1ed9 t\u1eadp trung: {sc.get('supplier_concentration', NOT_PROVIDED)}")
        r1.italic = True; r1.font.size = Pt(10)
        self._subheading(document, "2. Top 5 Kh\u00e1ch h\u00e0ng (Buyers)")
        t2 = self._table(document, ["#", "T\u00ean kh\u00e1ch h\u00e0ng", "Gi\u00e1 tr\u1ecb b\u00e1n", "T\u1ef7 l\u1ec7", "Terms", "S\u1ed1 n\u0103m QH"])
        for i, buyer in enumerate(sc.get("top_buyers", []), 1):
            self._row(t2, [str(i), str(buyer.get("name", NOT_PROVIDED)), _money(buyer.get("amount")), _percent(buyer.get("ratio")), str(buyer.get("payment_terms", NOT_PROVIDED)), str(buyer.get("relationship_years", NOT_PROVIDED))])
        p2 = document.add_paragraph()
        r2 = p2.add_run(f"T\u1ed5ng t\u1ef7 l\u1ec7 Top 5: {_percent(sc.get('top5_buyer_ratio'))}  |  M\u1ee9c \u0111\u1ed9 t\u1eadp trung: {sc.get('buyer_concentration', NOT_PROVIDED)}")
        r2.italic = True; r2.font.size = Pt(10)

    def _part5_statement_anomaly(self, document: Document, profile: Mapping[str, Any]) -> None:
        self._heading(document, "V. PH\u00c2N T\u00cdCH SAO K\u00ca NG\u00c2N H\u00c0NG & ANOMALY ENGINE")
        sa = profile.get("statement_anomaly_summary", {})
        self._subheading(document, "1. T\u1ed5ng quan d\u00f2ng ti\u1ec1n")
        t1 = self._table(document, ["Ch\u1ec9 ti\u00eau", "Gi\u00e1 tr\u1ecb"])
        self._row(t1, ["T\u1ed5ng inflow (12 th\u00e1ng)", _money(sa.get("total_inflow"))])
        self._row(t1, ["T\u1ed5ng outflow (12 th\u00e1ng)", _money(sa.get("total_outflow"))])
        self._row(t1, ["S\u1ed1 d\u01b0 b\u00ecnh qu\u00e2n th\u00e1ng", _money(sa.get("average_monthly_balance"))])
        self._row(t1, ["S\u1ed1 th\u00e1ng ph\u00e2n t\u00edch", str(sa.get("months_analyzed", NOT_PROVIDED))])
        self._subheading(document, "2. C\u1edd b\u00e1o \u0111\u1ed9ng anomaly (5 flags)")
        t2 = self._table(document, ["M\u00e3", "Tieu de", "Tr\u1ea1ng th\u00e1i", "M\u1ee9c \u0111\u1ed9", "S\u1ed1 l\u1ea7n", "T\u1ed5ng gi\u00e1 tr\u1ecb", "B\u1eb1ng ch\u1ee9ng"])
        for anomaly in sa.get("anomalies", []):
            status = "RED FLAG" if anomaly.get("triggered") else "PASS"
            self._row(t2, [str(anomaly.get("code", NOT_PROVIDED)), str(anomaly.get("title", NOT_PROVIDED)), status, str(anomaly.get("severity", "none")).upper(), str(anomaly.get("count", 0)), _money(anomaly.get("total_amount")), str(anomaly.get("evidence", NOT_PROVIDED))], alert=bool(anomaly.get("triggered")))

    def _part6_product_039(self, document: Document, profile: Mapping[str, Any]) -> None:
        self._heading(document, "VI. TH\u1ea8M \u0110\u1ecaNH \u0110I\u1ec0U KI\u1ec6N S\u1ea2N PH\u1ea8M \u0110\u1ea6U RA 039")
        product = profile.get("product_039_evaluation", {})
        table = self._table(document, ["#", "\u0110i\u1ec1u ki\u1ec7n", "K\u1ebft qu\u1ea3", "B\u1eb1ng ch\u1ee9ng"])
        for i, rule in enumerate(product.get("rules", []), 1):
            self._row(table, [str(i), str(rule.get("title")), str(rule.get("status")), str(rule.get("evidence"))], alert=rule.get("status") == "RED FLAG")
        decision = product.get("decision", NOT_PROVIDED)
        p = document.add_paragraph()
        r = p.add_run(f"K\u1ebft lu\u1eadn: {decision}")
        r.bold = True; r.font.size = Pt(11)
        if "T\u1eea ch\u1ed1i" in str(decision) or "T\u1eeA CH\u1ed0I" in str(decision):
            r.font.color.rgb = RGBColor(196, 18, 48)

    def _part7_swot(self, document: Document, profile: Mapping[str, Any]) -> None:
        self._heading(document, "VII. \u0110\u00c1NH GI\u00c1 T\u00cdN D\u1ee4NG 360\u00b0 (SWOT)")
        swot = profile.get("swot_analysis", {})
        t = self._table(document, ["\u0110i\u1ec3m m\u1ea1nh (Strengths)", "\u0110i\u1ec3m y\u1ebfu (Weaknesses)"])
        strengths = swot.get("strengths", [])
        weaknesses = swot.get("weaknesses", [])
        max_rows = max(len(strengths), len(weaknesses), 1)
        for i in range(max_rows):
            self._row(t, [str(strengths[i]) if i < len(strengths) else "", str(weaknesses[i]) if i < len(weaknesses) else ""])
        t2 = self._table(document, ["C\u01a1 h\u1ed9i (Opportunities)", "R\u1ee7i ro (Threats)"])
        opportunities = swot.get("opportunities", [])
        threats = swot.get("threats", [])
        max_rows2 = max(len(opportunities), len(threats), 1)
        for i in range(max_rows2):
            self._row(t2, [str(opportunities[i]) if i < len(opportunities) else "", str(threats[i]) if i < len(threats) else ""])

    def _part8_conclusion_signatures(self, document: Document, profile: Mapping[str, Any], source: Mapping[str, Any], warnings: list[str]) -> None:
        self._heading(document, "VIII. K\u1ebeT LU\u1eacN, \u0110I\u1ec0U KI\u1ec6N GI\u1ea2I NG\u00c2N & PH\u00ca DUY\u1ec6T")
        terms = source.get("credit_terms", {}) if isinstance(source.get("credit_terms"), Mapping) else {}
        pc = profile.get("product_039_evaluation", {}).get("calculations", {})
        proposed = terms.get("proposed_limit", pc.get("loan_request_amount"))
        if not terms:
            warnings.append("credit_terms was not supplied; non-limit credit terms are marked Ch\u01b0a cung c\u1ea5p.")
        self._subheading(document, "1. \u0110\u1ec1 xu\u1ea5t ph\u00ea duy\u1ec7t & \u0110i\u1ec1u ki\u1ec7n gi\u1ea3i ng\u00e2n (B43)")
        t1 = self._table(document, ["N\u1ed9i dung", "\u0110\u1ec1 xu\u1ea5t"])
        for label, value in [
            ("H\u1ea1n m\u1ee9c \u0111\u1ec1 xu\u1ea5t", _money(proposed)),
            ("Th\u1eddi h\u1ea1n c\u1ea5p HMTD", str(terms.get("tenor", "12 th\u00e1ng"))),
            ("Bi\u1ec7n ph\u00e1p b\u1ea3o \u0111\u1ea3m", str(terms.get("collateral", NOT_PROVIDED))),
            ("L\u00e3i su\u1ea5t", str(terms.get("interest_rate", NOT_PROVIDED))),
            ("\u0110i\u1ec1u ki\u1ec7n gi\u1ea3i ng\u00e2n", str(terms.get("disbursement_conditions", NOT_PROVIDED))),
            ("\u0110i\u1ec1u ki\u1ec7n q\u1ea3n l\u00fd sau GN", str(terms.get("post_disbursement_conditions", NOT_PROVIDED))),
        ]:
            self._row(t1, [label, value])
        outcome = profile.get("preliminary_decision", {}).get("outcome", "manual_review_required")
        p = document.add_paragraph()
        r = p.add_run(f"Khuy\u1ebfn ngh\u1ecb h\u1ec7 th\u1ed1ng: {outcome}")
        r.bold = True; r.font.size = Pt(11)
        p2 = document.add_paragraph()
        r2 = p2.add_run("K\u1ebft qu\u1ea3 l\u00e0 c\u00f4ng c\u1ee5 h\u1ed7 tr\u1ee3; c\u1ea5p c\u00f3 th\u1ea9m quy\u1ec1n th\u1ef1c hi\u1ec7n ph\u00ea duy\u1ec7t cu\u1ed1i c\u00f9ng sau khi ki\u1ec3m tra h\u1ed3 s\u01a1 g\u1ed1c.")
        r2.italic = True; r2.font.size = Pt(10)
        self._subheading(document, "2. Credit Covenants (Conditions Precedent & Subsequent)")
        covenants = profile.get("credit_covenants", {})
        self._subheading(document, "2a. \u0110i\u1ec1u ki\u1ec7n ti\u1ec1n \u0111i\u1ec1u (Conditions Precedent)")
        t_cp = self._table(document, ["#", "\u0110i\u1ec1u ki\u1ec7n"])
        for i, condition in enumerate(covenants.get("conditions_precedent", []), 1):
            self._row(t_cp, [str(i), str(condition)])
        self._subheading(document, "2b. \u0110i\u1ec1u ki\u1ec7n h\u1eadu \u0111i\u1ec1u (Conditions Subsequent)")
        t_cs = self._table(document, ["#", "\u0110i\u1ec1u ki\u1ec7n"])
        for i, condition in enumerate(covenants.get("conditions_subsequent", []), 1):
            self._row(t_cs, [str(i), str(condition)])
        self._subheading(document, "3. Khung 4 Ch\u1eef k\u00fd Ph\u00ea duy\u1ec7t (B55)")
        sig = self._table(document, [
            "Chuy\u00ean vi\u00ean QLKH (RM)\nL\u1eadp t\u1edd tr\u00ecnh",
            "Tr\u01b0\u1edfng ph\u00f2ng KHDN\nKi\u1ec3m so\u00e1t",
            "Gi\u00e1m \u0111\u1ed1c CN / Hub Tr\u01b0\u1edfng\nT\u00e1i th\u1ea9m \u0111\u1ecbnh & \u0110\u1ec1 xu\u1ea5t",
            "C\u1ea5p Ph\u00ea duy\u1ec7t TD RR H\u1ed9i s\u1edf\nPh\u00ea duy\u1ec7t",
        ])
        self._row(sig, ["\n\n\n\n(K\u00fd, ghi r\u00f5 h\u1ecd t\u00ean)\n\nNg\u00e0y:   /   /20___"] * 4)

    def _footer(self, document: Document) -> None:
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"T\u00e0i li\u1ec7u h\u1ed7 tr\u1ee3 th\u1ea9m \u0111\u1ecbnh n\u1ed9i b\u1ed9 | M-Credit 360 | M\u00e3 bi\u1ec3u m\u1eabu: {FORM_CODE}")
        run.font.size = Pt(8); run.italic = True

    def _table(self, document: Document, headers: list[str]) -> Any:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            _shade(cell, MSB_RED)
            _cell_border(cell)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        return table

    def _row(self, table: Any, values: list[str], alert: bool = False) -> None:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            _cell_border(cell)
            if alert:
                _shade(cell, ALERT_YELLOW)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(int(ALERT_TEXT[:2], 16), int(ALERT_TEXT[2:4], 16), int(ALERT_TEXT[4:], 16))
                    run.bold = True
            for run in cell.paragraphs[0].runs:
                if not run.font.size:
                    run.font.size = Pt(10)


class _RequestHandler(BaseHTTPRequestHandler):
    builder = CreditMemoBuilder()

    def do_GET(self) -> None:  # noqa: N802
        if self.path.rstrip("/") == "/health":
            self._send(HTTPStatus.OK, {"status": "ok", "agent": "credit-memo-builder"})
        else:
            self._send(HTTPStatus.NOT_FOUND, {"error": "Not found"})

    def do_POST(self) -> None:  # noqa: N802
        if self.path.rstrip("/") != "/build-memo-docx":
            self._send(HTTPStatus.NOT_FOUND, {"error": "Not found"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if not 0 < length <= 10_000_000:
                raise MemoBuilderError("Request body must be between 1 byte and 10 MB.")
            request = json.loads(self.rfile.read(length).decode("utf-8"))
            self._send(HTTPStatus.OK, self.builder.build(request))
        except (MemoBuilderError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            self._send(HTTPStatus.BAD_REQUEST, {"error": str(exc)})
        except Exception:
            LOGGER.exception("Credit memo build failed")
            self._send(HTTPStatus.INTERNAL_SERVER_ERROR, {"error": "Internal memo build error"})

    def _send(self, status: HTTPStatus, body: Mapping[str, Any]) -> None:
        encoded = json.dumps(body, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def log_message(self, format: str, *args: Any) -> None:
        LOGGER.info("%s - %s", self.address_string(), format % args)


def run_server() -> None:
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
    port = int(os.getenv("PORT", "8080"))
    server = ThreadingHTTPServer(("0.0.0.0", port), _RequestHandler)
    LOGGER.info("Credit Memo Builder listening on port %s", port)
    server.serve_forever()


if __name__ == "__main__":
    run_server()
