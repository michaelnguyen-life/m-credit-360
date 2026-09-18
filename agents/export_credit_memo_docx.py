"""Export a Markdown credit memo to an editable Word document."""

from __future__ import annotations

import sys
import re
from pathlib import Path


def export(markdown_path: Path, output_path: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("Install python-docx to export Word documents.") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        headers, *rows = table_rows
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        table_rows.clear()

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            values = [value.strip() for value in stripped.strip("|").split("|")]
            if all(value.replace(":", "").replace("-", "") == "" for value in values):
                continue
            table_rows.append(values)
            continue
        flush_table()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
        else:
            document.add_paragraph(stripped.replace("**", ""))
    flush_table()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python export_credit_memo_docx.py <memo.md> <output.docx>")
    export(Path(sys.argv[1]), Path(sys.argv[2]))
