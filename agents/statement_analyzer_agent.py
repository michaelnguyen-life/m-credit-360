"""Bank Statement Analyzer & Anomaly Detector for M-Credit 360."""

from __future__ import annotations

import csv
import json
import logging
import os
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Iterable, Mapping


LOGGER = logging.getLogger(__name__)
VND = Decimal("1000000")
DEFAULT_RULES: dict[str, Any] = {
    "large_cash_withdrawal": Decimal("100") * VND,
    "cash_withdrawal_share": Decimal("0.30"),
    "rapid_movement_window_hours": 24,
    "rapid_movement_share": Decimal("0.90"),
    "round_amount_unit": Decimal("50") * VND,
    "round_amount_min_count": 3,
    "round_amount_window_days": 30,
    "after_hours_start": 23,
    "after_hours_end": 5,
    "loan_keywords": ("vay", "tra no", "dao han", "borrow", "loan repayment"),
    "internal_transfer_keywords": ("chuyen noi bo", "noi bo", "tai khoan cua toi", "internal transfer", "own account"),
    "sensitive_keywords": ("vay", "tra no", "dao han", "cam do", "tai xiu", "crypto", "tien ao"),
}

COLUMN_ALIASES = {
    "occurred_at": ("ngay giao dich", "ngay", "transaction date", "date", "thoi gian giao dich"),
    "inflow": ("ghi co", "so tien ghi co", "credit", "inflow", "tien vao"),
    "outflow": ("ghi no", "so tien ghi no", "debit", "outflow", "tien ra"),
    "balance": ("so du cuoi", "so du", "closing balance", "balance"),
    "description": ("noi dung giao dich", "noi dung", "description", "transaction content", "dien giai"),
}


class StatementError(ValueError):
    """A request or source document cannot be safely analyzed."""


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip().replace(" ", "")
    if not text or text in {"-", "--", "n/a"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    # A lone comma between digit groups is a decimal separator; otherwise both
    # separators are treated as thousands separators for Vietnamese statements.
    if text.count(",") == 1 and "." not in text and len(text.rsplit(",", 1)[1]) <= 2:
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "").replace(".", "")
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return -number if negative else number


def _number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _fold(text: Any) -> str:
    decomposed = unicodedata.normalize("NFD", str(text).lower())
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn").replace("đ", "d")


def _contains(description: str, keywords: Iterable[str]) -> bool:
    normalized = _fold(description)
    return any(_fold(keyword) in normalized for keyword in keywords)


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = str(value or "").strip()
    for layout in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, layout)
        except ValueError:
            pass
    return None


@dataclass(frozen=True)
class Transaction:
    transaction_id: str
    occurred_at: datetime
    inflow: Decimal
    outflow: Decimal
    balance: Decimal | None
    description: str
    source_row: int
    time_known: bool

    def as_dict(self) -> dict[str, Any]:
        return {"transaction_id": self.transaction_id, "occurred_at": self.occurred_at.isoformat(sep=" "), "inflow": _number(self.inflow), "outflow": _number(self.outflow), "balance": _number(self.balance), "description": self.description, "source_row": self.source_row}


@dataclass
class StatementAnalyzer:
    rules: dict[str, Any] = field(default_factory=lambda: dict(DEFAULT_RULES))

    def analyze(self, payload: Mapping[str, Any]) -> dict[str, Any]:
        if not isinstance(payload, Mapping):
            raise StatementError("Payload must be a JSON object.")
        warnings: list[str] = []
        rules = self._effective_rules(payload.get("rules"), warnings)
        transactions = self._load_transactions(payload, warnings)
        if not transactions:
            raise StatementError("No usable statement transactions were found.")
        transactions.sort(key=lambda item: item.occurred_at)
        monthly = self._monthly_summary(transactions)
        excluded, eligible_income = self._eligible_income(transactions, rules)
        anomalies = self._detect_anomalies(transactions, monthly, rules)
        report = {
            "schema_version": "1.0",
            "assessment_id": str(payload.get("assessment_id") or f"STMT-{datetime.now(timezone.utc):%Y%m%d%H%M%S}"),
            "assessed_at": datetime.now(timezone.utc).isoformat(),
            "status": "completed_with_warnings" if warnings else "completed",
            "account": dict(payload.get("account", {})) if isinstance(payload.get("account"), Mapping) else {},
            "source": {"type": payload.get("source_type", "transactions"), "transaction_count": len(transactions)},
            "monthly_cash_flow": monthly,
            "net_eligible_income": eligible_income,
            "excluded_transactions": excluded,
            "suspicious_transactions": anomalies,
            "risk_summary": self._risk_summary(anomalies),
            "rules_applied": self._serializable_rules(rules),
            "warnings": warnings,
            "errors": [],
        }
        report["statement_assessment_report"] = self.render_markdown(report)
        return report

    def analyze_file(self, path: str | Path, **metadata: Any) -> dict[str, Any]:
        source = Path(path)
        if not source.is_file():
            raise StatementError(f"Statement file does not exist: {source}")
        suffix = source.suffix.lower()
        if suffix == ".csv":
            rows = self._read_csv(source)
        elif suffix == ".xlsx":
            rows = self._read_xlsx(source)
        elif suffix == ".pdf":
            rows = self._read_pdf(source)
        else:
            raise StatementError("Supported statement formats are .csv, .xlsx and text-based .pdf.")
        return self.analyze({**metadata, "transactions": rows, "source_type": suffix[1:]})

    def write_outputs(self, report: Mapping[str, Any], output_dir: str | Path = "output_memos") -> tuple[Path, Path]:
        directory = Path(output_dir)
        directory.mkdir(parents=True, exist_ok=True)
        identifier = re.sub(r"[^A-Za-z0-9_.-]", "_", str(report["assessment_id"]))
        json_path = directory / f"{identifier}.statement_assessment.json"
        markdown_path = directory / f"{identifier}.statement_assessment.md"
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        markdown_path.write_text(str(report["statement_assessment_report"]), encoding="utf-8")
        return json_path, markdown_path

    def _load_transactions(self, payload: Mapping[str, Any], warnings: list[str]) -> list[Transaction]:
        if "file_path" in payload:
            raise StatementError("Use analyze_file() for file_path input, or submit normalized transactions to the HTTP endpoint.")
        raw = payload.get("transactions")
        if not isinstance(raw, list):
            raise StatementError("transactions must be a list of transaction objects.")
        result: list[Transaction] = []
        for index, row in enumerate(raw, start=1):
            if not isinstance(row, Mapping):
                warnings.append(f"Ignored non-object transaction at row {index}.")
                continue
            normalized = self._normalize_row(row)
            when = _parse_datetime(normalized.get("occurred_at"))
            inflow = _decimal(normalized.get("inflow")) or Decimal("0")
            outflow = _decimal(normalized.get("outflow")) or Decimal("0")
            if when is None or (inflow <= 0 and outflow <= 0):
                warnings.append(f"Ignored invalid transaction at row {index}: date and a positive inflow/outflow are required.")
                continue
            result.append(Transaction(str(normalized.get("transaction_id") or f"TX-{index:05d}"), when, abs(inflow), abs(outflow), _decimal(normalized.get("balance")), str(normalized.get("description") or ""), index, bool(re.search(r"\d{1,2}:\d{2}", str(normalized.get("occurred_at"))))))
        return result

    @staticmethod
    def _normalize_row(row: Mapping[str, Any]) -> dict[str, Any]:
        normalized = dict(row)
        aliases = {_fold(key): key for key in row}
        for target, options in COLUMN_ALIASES.items():
            if target in normalized:
                continue
            for alias in options:
                source_key = aliases.get(_fold(alias))
                if source_key is not None:
                    normalized[target] = row[source_key]
                    break
        return normalized

    def _monthly_summary(self, transactions: list[Transaction]) -> list[dict[str, Any]]:
        buckets: dict[str, list[Transaction]] = defaultdict(list)
        for transaction in transactions:
            buckets[transaction.occurred_at.strftime("%Y-%m")].append(transaction)
        result = []
        for month, items in sorted(buckets.items()):
            balances = [item.balance for item in items if item.balance is not None]
            inflow = sum((item.inflow for item in items), Decimal("0"))
            outflow = sum((item.outflow for item in items), Decimal("0"))
            result.append({"month": month, "total_inflow": _number(inflow), "total_outflow": _number(outflow), "net_cash_flow": _number(inflow - outflow), "transaction_count": len(items), "average_balance": _number(sum(balances, Decimal("0")) / len(balances)) if balances else None, "minimum_balance": _number(min(balances)) if balances else None, "maximum_balance": _number(max(balances)) if balances else None})
        return result

    def _eligible_income(self, transactions: list[Transaction], rules: Mapping[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        excluded: list[dict[str, Any]] = []
        eligible_inflow = Decimal("0")
        eligible_outflow = Decimal("0")
        for transaction in transactions:
            loan = _contains(transaction.description, rules["loan_keywords"])
            internal = _contains(transaction.description, rules["internal_transfer_keywords"])
            if transaction.inflow and (loan or internal):
                excluded.append({**transaction.as_dict(), "reason": "loan_or_debt_transaction" if loan else "internal_transfer"})
            else:
                eligible_inflow += transaction.inflow
            if transaction.outflow and not internal:
                eligible_outflow += transaction.outflow
        return excluded, {"eligible_inflow": _number(eligible_inflow), "eligible_outflow": _number(eligible_outflow), "net_eligible_income": _number(eligible_inflow - eligible_outflow), "methodology": "Eligible inflow excludes loan/debt and internal-transfer descriptions; eligible outflow excludes identified internal transfers."}

    def _detect_anomalies(self, transactions: list[Transaction], monthly: list[Mapping[str, Any]], rules: Mapping[str, Any]) -> list[dict[str, Any]]:
        findings: list[dict[str, Any]] = []
        def add(transaction: Transaction, tag: str, rule: str, evidence: str, related: str | None = None) -> None:
            findings.append({"transaction": transaction.as_dict(), "risk_tag": tag, "rule": rule, "evidence": evidence, "related_transaction_id": related})
        monthly_outflow = {item["month"]: Decimal(str(item["total_outflow"])) for item in monthly}
        cash_items = [item for item in transactions if item.outflow and _contains(item.description, ("rut tien mat", "cash withdrawal", "atm"))]
        for item in cash_items:
            total = monthly_outflow[item.occurred_at.strftime("%Y-%m")]
            share = item.outflow / total if total else Decimal("0")
            if item.outflow > rules["large_cash_withdrawal"]:
                add(item, "HIGH", "CASH_WITHDRAWAL", f"Cash withdrawal {_number(item.outflow)} exceeds {_number(rules['large_cash_withdrawal'])}.")
            elif share > rules["cash_withdrawal_share"]:
                add(item, "MEDIUM", "CASH_WITHDRAWAL", f"Cash withdrawal is {_number(share * 100)}% of monthly outflow.")
        self._rapid_movement(transactions, rules, add)
        self._repeated_round_amounts(transactions, rules, add)
        for item in transactions:
            if item.time_known and (item.occurred_at.hour >= rules["after_hours_start"] or item.occurred_at.hour < rules["after_hours_end"]):
                add(item, "MEDIUM", "AFTER_HOURS", f"Transaction timestamp {item.occurred_at:%H:%M} falls in the 23:00-05:00 risk window.")
            keyword = next((word for word in rules["sensitive_keywords"] if _fold(word) in _fold(item.description)), None)
            if keyword:
                tag = "CRITICAL" if _fold(keyword) in {"cam do", "tai xiu"} else "HIGH"
                add(item, tag, "SENSITIVE_KEYWORD", f"Sensitive keyword detected: {keyword}.")
        return findings

    @staticmethod
    def _rapid_movement(transactions: list[Transaction], rules: Mapping[str, Any], add: Any) -> None:
        for index, incoming in enumerate(transactions):
            if incoming.inflow <= 0:
                continue
            outgoing = Decimal("0")
            for candidate in transactions[index + 1:]:
                if candidate.occurred_at - incoming.occurred_at > timedelta(hours=rules["rapid_movement_window_hours"]):
                    break
                outgoing += candidate.outflow
                if outgoing >= incoming.inflow * rules["rapid_movement_share"]:
                    add(incoming, "HIGH", "RAPID_FUNDS_MOVEMENT", f"Outgoing transactions reached {_number(outgoing / incoming.inflow * 100)}% of inflow within {rules['rapid_movement_window_hours']} hours.", candidate.transaction_id)
                    break

    @staticmethod
    def _repeated_round_amounts(transactions: list[Transaction], rules: Mapping[str, Any], add: Any) -> None:
        unit = rules["round_amount_unit"]
        candidates = [item for item in transactions if (item.inflow or item.outflow) >= unit and (item.inflow or item.outflow) % unit == 0]
        for item in candidates:
            nearby = [other for other in candidates if 0 <= (other.occurred_at - item.occurred_at).days <= rules["round_amount_window_days"]]
            if len(nearby) >= rules["round_amount_min_count"]:
                add(item, "MEDIUM", "REPEATED_ROUND_AMOUNT", f"{len(nearby)} round-amount transactions of at least {_number(unit)} occur within {rules['round_amount_window_days']} days.")

    def _effective_rules(self, custom: Any, warnings: list[str]) -> dict[str, Any]:
        effective = dict(self.rules)
        if custom is None:
            return effective
        if not isinstance(custom, Mapping):
            warnings.append("Ignored rules because it is not an object.")
            return effective
        numeric = {"large_cash_withdrawal", "cash_withdrawal_share", "rapid_movement_share", "round_amount_unit"}
        integers = {"rapid_movement_window_hours", "round_amount_min_count", "round_amount_window_days", "after_hours_start", "after_hours_end"}
        for key, value in custom.items():
            if key in numeric:
                parsed = _decimal(value)
                if parsed is not None and parsed >= 0:
                    effective[key] = parsed
                else:
                    warnings.append(f"Ignored invalid rule {key}.")
            elif key in integers and isinstance(value, int) and value >= 0:
                effective[key] = value
            elif key.endswith("_keywords") and isinstance(value, list) and all(isinstance(word, str) for word in value):
                effective[key] = tuple(value)
            else:
                warnings.append(f"Ignored unsupported or invalid rule {key}.")
        return effective

    @staticmethod
    def _risk_summary(findings: list[Mapping[str, Any]]) -> dict[str, Any]:
        counts = {tag: sum(item["risk_tag"] == tag for item in findings) for tag in ("CRITICAL", "HIGH", "MEDIUM")}
        return {"total_suspicious_transactions": len(findings), "by_risk_tag": counts, "manual_review_required": counts["CRITICAL"] > 0 or counts["HIGH"] >= 2}

    @staticmethod
    def _serializable_rules(rules: Mapping[str, Any]) -> dict[str, Any]:
        return {key: _number(value) if isinstance(value, Decimal) else list(value) if isinstance(value, tuple) else value for key, value in rules.items()}

    @staticmethod
    def render_markdown(report: Mapping[str, Any]) -> str:
        lines = ["# Statement Assessment Report", f"**Assessment ID:** {report['assessment_id']}  ", f"**Transactions analyzed:** {report['source']['transaction_count']}", "", "## Monthly Cash Flow", "| Month | Inflow | Outflow | Net Flow | Avg. Balance | Min / Max Balance |", "|---|---:|---:|---:|---:|---:|"]
        for month in report["monthly_cash_flow"]:
            lines.append(f"| {month['month']} | {month['total_inflow']} | {month['total_outflow']} | {month['net_cash_flow']} | {month['average_balance']} | {month['minimum_balance']} / {month['maximum_balance']} |")
        income = report["net_eligible_income"]
        lines.extend(["", "## Net Eligible Income", f"- Eligible inflow: {income['eligible_inflow']}", f"- Eligible outflow: {income['eligible_outflow']}", f"- **Net Eligible Income: {income['net_eligible_income']}**", "", "## Suspicious Transactions"])
        if report["suspicious_transactions"]:
            for finding in report["suspicious_transactions"]:
                tx = finding["transaction"]
                lines.append(f"- [{finding['risk_tag']}] {finding['rule']} | {tx['occurred_at']} | {tx['transaction_id']} | {finding['evidence']}")
        else:
            lines.append("- No rule-based anomalies detected.")
        lines.extend(["", "## Review Note", "This report supports credit assessment only. A credit officer must validate source statements and transaction context before making any credit decision."])
        return "\n".join(lines) + "\n"

    @staticmethod
    def _read_csv(path: Path) -> list[dict[str, Any]]:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="cp1258")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(text.splitlines(), dialect=dialect))

    @staticmethod
    def _read_xlsx(path: Path) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook  # type: ignore[import-not-found]
        except ImportError as exc:
            raise StatementError("XLSX support requires installing openpyxl.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "") for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]

    @staticmethod
    def _read_pdf(path: Path) -> list[dict[str, Any]]:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except ImportError as exc:
            raise StatementError("PDF support requires installing pypdf.") from exc
        try:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        except Exception as exc:
            raise StatementError(f"Could not read PDF: {exc}") from exc
        if not text.strip():
            raise StatementError("PDF has no extractable text; OCR is required for scanned statements.")
        rows = []
        date_pattern = re.compile(r"^(\d{2}/\d{2}/\d{4}(?:\s+\d{2}:\d{2}(?::\d{2})?)?)\s+(.+?)\s+([\d.,()-]+)\s+([\d.,()-]+)(?:\s+([\d.,()-]+))?$")
        for line in text.splitlines():
            matched = date_pattern.match(line.strip())
            if matched:
                date, description, inflow, outflow, balance = matched.groups()
                rows.append({"occurred_at": date, "description": description, "inflow": inflow, "outflow": outflow, "balance": balance})
        if not rows:
            raise StatementError("Could not reliably identify transaction rows in this PDF. Provide CSV/XLSX or bank-specific parser.")
        return rows


class _RequestHandler(BaseHTTPRequestHandler):
    agent = StatementAnalyzer()

    def do_GET(self) -> None:  # noqa: N802
        self._send(HTTPStatus.OK, {"status": "ok", "agent": "statement-analyzer"}) if self.path.rstrip("/") == "/health" else self._send(HTTPStatus.NOT_FOUND, {"error": "Not found"})

    def do_POST(self) -> None:  # noqa: N802
        if self.path.rstrip("/") != "/analyze-statement":
            self._send(HTTPStatus.NOT_FOUND, {"error": "Not found"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if not 0 < length <= 10_000_000:
                raise StatementError("Request body must be between 1 byte and 10 MB.")
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            if not isinstance(payload, Mapping):
                raise StatementError("Payload must be a JSON object.")
            report = self.agent.analyze(payload)
            self._send(HTTPStatus.OK, report)
        except (StatementError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            self._send(HTTPStatus.BAD_REQUEST, {"error": str(exc)})
        except Exception:
            LOGGER.exception("Unexpected statement analysis error")
            self._send(HTTPStatus.INTERNAL_SERVER_ERROR, {"error": "Internal statement analysis error"})

    def _send(self, status: HTTPStatus, body: Mapping[str, Any]) -> None:
        encoded = json.dumps(body, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def log_message(self, format: str, *args: Any) -> None:
        LOGGER.info("%s - %s", self.address_string(), format % args)


def run_server() -> None:
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
    port = int(os.getenv("PORT", "8080"))
    server = ThreadingHTTPServer(("0.0.0.0", port), _RequestHandler)
    LOGGER.info("Statement Analyzer listening on port %s", port)
    server.serve_forever()


if __name__ == "__main__":
    run_server()
