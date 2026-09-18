import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
from credit_memo_builder_agent import CreditMemoBuilder, FORM_CODE


class CreditMemoBuilderTests(unittest.TestCase):
    """Verify the MB02a/QT.RR.037 credit memo builder produces a valid DOCX."""

    @classmethod
    def setUpClass(cls):
        cls.payload_path = (
            Path(__file__).parent.parent
            / "data_test"
            / "1. CTY DAU TU GROUP (MOCK AN DANH)"
            / "eb_credit_payload.json"
        )
        cls.payload = json.loads(cls.payload_path.read_text(encoding="utf-8"))

    def _build(self, directory):
        return CreditMemoBuilder(directory).build({"eb_payload": self.payload})

    def test_builds_valid_docx_with_correct_filename(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            output = Path(result["file_path"])
            self.assertTrue(output.is_file())
            self.assertEqual(
                output.name,
                "TO_TRINH_TIN_DUNG_MB02a_0318999888_2025.docx",
            )

    def test_docx_can_be_reopened(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            self.assertGreater(len(document.paragraphs), 0)
            self.assertGreaterEqual(len(document.tables), 15)

    def test_header_contains_form_code_and_bank_name(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("MSB", text)
            self.assertIn(FORM_CODE, text)

    def _full_text(self, document):
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    def test_all_eight_sections_present(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("I. TH\u00d4NG TIN KH\u00c1CH H\u00c0NG & PH\u00c1P L\u00dd", text)
            self.assertIn("II. \u0110\u1ec0 XU\u1ea4T C\u1ea4P T\u00cdN D\u1ee4NG", text)
            self.assertIn("III. N\u1ed8I DUNG TH\u1ea8M \u0110\u1ecaNH & T\u00c0I CH\u00cdNH BCTC 3 K\u1ef2", text)
            self.assertIn("IV. PH\u00c2N T\u00cdCH CHU\u1ed4I CUNG \u1ee8NG", text)
            self.assertIn("V. PH\u00c2N T\u00cdCH SAO K\u00ca NG\u00c2N H\u00c0NG", text)
            self.assertIn("VI. TH\u1ea8M \u0110\u1ecaNH \u0110I\u1ec0U KI\u1ec6N S\u1ea2N PH\u1ea8M \u0110\u1ea6U RA 039", text)
            self.assertIn("VII. \u0110\u00c1NH GI\u00c1 T\u00cdN D\u1ee4NG 360\u00b0", text)
            self.assertIn("VIII. K\u1ebeT LU\u1eacN", text)

    def test_supply_chain_tables_present(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = self._full_text(document)
            self.assertIn("Top 5 Nh\u00e0 cung c\u1ea5p", text)
            self.assertIn("Top 5 Kh\u00e1ch h\u00e0ng", text)
            self.assertIn("CTCP VLXD S\u00c0I G\u00d2N", text)

    def test_anomaly_engine_present(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = self._full_text(document)
            self.assertIn("ANOMALY ENGINE", text)
            self.assertIn("ANM01_LARGE_CASH_WITHDRAWAL", text)
            self.assertIn("ANM05_SENSITIVE_KEYWORDS", text)

    def test_swot_tables_present(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = self._full_text(document)
            self.assertIn("\u0110i\u1ec3m m\u1ea1nh (Strengths)", text)
            self.assertIn("R\u1ee7i ro (Threats)", text)

    def test_credit_covenants_present(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("Conditions Precedent", text)
            self.assertIn("Conditions Subsequent", text)

    def test_part1_has_four_column_kyc_table(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            kyc_tables = [
                t for t in document.tables
                if len(t.columns) == 4 and t.rows[0].cells[0].text.strip() == "H\u1ea1ng m\u1ee5c"
            ]
            self.assertGreaterEqual(len(kyc_tables), 1)

    def test_part5_has_four_column_signature_table(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            document = Document(result["file_path"])
            sig_tables = [
                t for t in document.tables
                if len(t.columns) == 4 and "QLKH" in t.rows[0].cells[0].text
            ]
            self.assertGreaterEqual(len(sig_tables), 1)

    def test_build_accepts_credit_profile_directly(self):
        with tempfile.TemporaryDirectory() as d:
            profile = CreditMemoBuilder(d).credit_agent.assess(self.payload)
            result = CreditMemoBuilder(d).build({"credit_profile": profile})
            self.assertTrue(Path(result["file_path"]).is_file())

    def test_build_returns_expected_metadata(self):
        with tempfile.TemporaryDirectory() as d:
            result = self._build(d)
            self.assertIn(result["status"], ("completed", "completed_with_warnings"))
            self.assertEqual(result["assessment_id"], "EB-ALPHA-2025-001")
            self.assertEqual(
                result["product_039_decision"],
                "\u0110\u1ee6 \u0110I\u1ec0U KI\u1ec6N S\u1ea2N PH\u1ea8M 039",
            )


if __name__ == "__main__":
    unittest.main()
