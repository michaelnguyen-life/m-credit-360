import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Tiêu đề
title = doc.add_heading('BẢN ĐỒ CẤU TRÚC GIAO DIỆN M-CREDIT 360', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Tài liệu này giúp Sếp Michael Nguyên nắm rõ cấu trúc các khối (Blocks) trên trang chủ index.html để có thể tự thay đổi nội dung (Content) mà không làm vỡ giao diện (Layout) hay làm hỏng Logic (Javascript).')

doc.add_heading('I. NGUYÊN TẮC BẤT DI BẤT DỊCH TỪ ĐẶC VỤ', level=1)
p = doc.add_paragraph()
p.add_run('1. Chỉ sửa nội dung TEXT (chữ tiếng Việt) nằm giữa các thẻ HTML (ví dụ: >Chữ cần sửa<).\n')
p.add_run('2. TUYỆT ĐỐI KHÔNG sửa các id="...", class="...", hoặc onclick="..." vì đây là các khớp nối thần kinh của Agent.\n')
p.add_run('3. TUYỆT ĐỐI KHÔNG dùng Powershell (Get-Content / Set-Content) để sửa file index.html vì sẽ bị lỗi font Tiếng Việt (Mojibake). Hãy mở bằng VSCode hoặc Notepad++.')

doc.add_heading('II. CẤU TRÚC 3 CỘT (GRID 4-5-3)', level=1)
doc.add_paragraph('Giao diện Khách hàng Doanh nghiệp (EB) được chia thành 12 phần (Grid 12), chia làm 3 cột chính:')

# Cột 1
doc.add_heading('1. CỘT TRÁI (col-span-4): Form Nhập liệu', level=2)
doc.add_paragraph('Bao gồm các nút điều khiển chính và Form nhập số liệu BCTC.')
doc.add_paragraph('- Các nút: "+ Khách hàng mới", "Upload hồ sơ", "Alpha Group", "Beta Corp", "CHẠY THẨM ĐỊNH AI 360°".\n- Vị trí sửa an toàn: Tên các nút bấm (lưu ý không chạm vào onclick).')

# Cột 2
doc.add_heading('2. CỘT GIỮA (col-span-5): Kết quả & Cảnh báo', level=2)
doc.add_paragraph('Đây là khu vực hiển thị kết quả động từ Server. Các nội dung tĩnh có thể sửa:')
doc.add_paragraph('- Tiêu đề: "KẾT QUẢ THẨM ĐỊNH SƠ BỘ MSB"\n- Tiêu đề 4 KPIs: Vốn lưu động ròng, Hệ số trả nợ (DSCR), Bù đắp lãi vay (ICR), Tài trợ QĐ 039.\n- Khối "Trinh sát Dữ liệu (OSINT 360°)": Đã chốt 2 mục Nợ Thuế và Mua Sắm Công.\n- Khối "Cảnh báo rủi ro (Red Flags)"\n- Khối "Cơ hội Bán chéo (Cross-sell Alert)": Nội dung bên trong khối này được Render TỰ ĐỘNG bằng Javascript từ file eb_credit_agent.py, không thể sửa trực tiếp bằng HTML tĩnh.')

# Cột 3
doc.add_heading('3. CỘT PHẢI (col-span-3): Danh mục Tài liệu', level=2)
doc.add_paragraph('Khu vực chứa danh sách các file PDF, Excel đã tải lên.')
doc.add_paragraph('- Box 1: Hồ sơ khách hàng (Hiển thị các file đã tải).\n- Ghi chú: Block "M-CREDIT AI Insight" chốt cứng đã BỊ XÓA VĨNH VIỄN theo lệnh của Sếp để đảm bảo tính logic cho KH mới.')

doc.add_heading('III. CẤU TRÚC JAVASCRIPT (NÃO BỘ)', level=1)
doc.add_paragraph('Nằm ở cuối file index.html, bên trong thẻ <script>. Đây là não bộ điều khiển:')
doc.add_paragraph('1. runEBAssessment(): Kích hoạt API thẩm định.\n2. saveEBProfile(): Lưu hồ sơ vào DB.\n3. loadEBProfile(mst): Tự động gọi lại hồ sơ khi gõ MST.\n4. openNewCustomerModal(): Xóa trắng Form.')

doc.save('Cau_truc_M_CREDIT_360.docx')
